from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from fastapi.responses import StreamingResponse
from typing import List, Optional
from datetime import datetime
import io
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, PatternFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from database import db
from models.telematics import TelematicsRecord, TelematicsCreate
from models.user import User
from core.security import get_current_user

router = APIRouter(prefix="/telematics", tags=["telematics"])
_proj = {"_id": 0}

# ── Column specification ───────────────────────────────────────────────────────
COLUMNS = [
    {
        "field":       "registration_number",
        "header":      "Vehicle_Registration",
        "description": "Vehicle registration number (must match a vehicle in the system)",
        "example":     "ABC1234",
        "width":       24,
        "type":        "str",
        "required":    True,
    },
    {
        "field":       "customer_name",
        "header":      "Customer_Name",
        "description": "Full name of the insured (for reference only)",
        "example":     "John Smith",
        "width":       22,
        "type":        "str",
        "required":    False,
    },
    {
        "field":       "period_month",
        "header":      "Period_Month",
        "description": "Data collection month in YYYY-MM format",
        "example":     datetime.utcnow().strftime("%Y-%m"),
        "width":       16,
        "type":        "str",
        "required":    True,
    },
    {
        "field":       "avg_monthly_mileage_km",
        "header":      "Avg_Monthly_Mileage_km",
        "description": "Average km driven per month (50 – 5000)",
        "example":     1200,
        "width":       26,
        "type":        "float",
        "min":         50,
        "max":         5000,
        "required":    True,
    },
    {
        "field":       "hard_braking_events_monthly",
        "header":      "Hard_Braking_Events",
        "description": "Number of hard-braking events recorded this month (0 – 20)",
        "example":     3,
        "width":       24,
        "type":        "int",
        "min":         0,
        "max":         20,
        "required":    True,
    },
    {
        "field":       "speeding_incidents_monthly",
        "header":      "Speeding_Incidents",
        "description": "Number of speeding incidents recorded this month (0 – 10)",
        "example":     1,
        "width":       22,
        "type":        "int",
        "min":         0,
        "max":         10,
        "required":    True,
    },
    {
        "field":       "night_driving_pct",
        "header":      "Night_Driving_Pct",
        "description": "Percentage of total driving done between 22:00 – 05:00 (0 – 100)",
        "example":     15.0,
        "width":       22,
        "type":        "float",
        "min":         0,
        "max":         100,
        "required":    True,
    },
    {
        "field":       "notes",
        "header":      "Notes",
        "description": "Optional notes or comments",
        "example":     "",
        "width":       30,
        "type":        "str",
        "required":    False,
    },
]

# Brand colours
BLUE_DARK   = "1E3A8A"
BLUE_MID    = "2563EB"
BLUE_LIGHT  = "DBEAFE"
BLUE_PALE   = "EFF6FF"
ORANGE      = "EA580C"
WHITE       = "FFFFFF"
GRAY_LIGHT  = "F8FAFC"
GRAY_BORDER = "CBD5E1"


def _border(style="thin"):
    s = Side(style=style, color=GRAY_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)


def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


# ── Template builder ───────────────────────────────────────────────────────────

def build_template_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()

    # ── Sheet 1: Data entry ────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Telematics Data"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    # Row 1: title banner
    ws.merge_cells(f"A1:{get_column_letter(len(COLUMNS))}1")
    title_cell = ws["A1"]
    title_cell.value = "AutoRisk Premium Optimizer — Telematics / Behavioural Data Upload"
    title_cell.font = Font(bold=True, size=13, color=WHITE, name="Calibri")
    title_cell.fill = _fill(BLUE_DARK)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Row 2: column headers
    for col_idx, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx)
        cell.value = col["header"]
        req = " *" if col["required"] else ""
        cell.value = col["header"] + req
        cell.font = Font(bold=True, size=10, color=WHITE, name="Calibri")
        cell.fill = _fill(BLUE_MID)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _border()
        ws.column_dimensions[get_column_letter(col_idx)].width = col["width"]
    ws.row_dimensions[2].height = 36

    # Row 3: example / sample data (light blue)
    for col_idx, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = col["example"]
        cell.font = Font(italic=True, size=10, color="475569", name="Calibri")
        cell.fill = _fill(BLUE_LIGHT)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _border()
    ws.row_dimensions[3].height = 20

    # Rows 4-53: data entry rows (50 blank rows)
    for row in range(4, 54):
        for col_idx in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.fill = _fill(GRAY_LIGHT if row % 2 == 0 else WHITE)
            cell.border = _border()
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 18

    # Data validations for numeric columns
    for col_idx, col in enumerate(COLUMNS, start=1):
        if col["type"] in ("int", "float") and "min" in col:
            dv = DataValidation(
                type="decimal",
                operator="between",
                formula1=str(col["min"]),
                formula2=str(col["max"]),
                showErrorMessage=True,
                errorTitle="Out of range",
                error=f"{col['header']} must be between {col['min']} and {col['max']}.",
                showInputMessage=True,
                promptTitle=col["header"],
                prompt=col["description"],
            )
            col_letter = get_column_letter(col_idx)
            dv.sqref = f"{col_letter}4:{col_letter}53"
            ws.add_data_validation(dv)

    # Print settings
    ws.print_title_rows = "1:2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1

    # ── Sheet 2: Instructions ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("Instructions")
    ws2.sheet_view.showGridLines = False
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 65

    # Title
    ws2.merge_cells("A1:B1")
    t = ws2["A1"]
    t.value = "AutoRisk — Telematics Upload Instructions"
    t.font = Font(bold=True, size=13, color=WHITE, name="Calibri")
    t.fill = _fill(ORANGE)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    # Sub-header
    ws2.merge_cells("A2:B2")
    s = ws2["A2"]
    s.value = "Complete the 'Telematics Data' sheet and upload it through the AutoRisk system."
    s.font = Font(size=10, color="475569", italic=True, name="Calibri")
    s.alignment = Alignment(horizontal="center")
    ws2.row_dimensions[2].height = 20

    # Column descriptions
    ws2.cell(row=3, column=1).value = "Column"
    ws2.cell(row=3, column=2).value = "Description & Accepted Values"
    for c in [ws2.cell(row=3, column=1), ws2.cell(row=3, column=2)]:
        c.font = Font(bold=True, size=10, color=WHITE, name="Calibri")
        c.fill = _fill(BLUE_DARK)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _border()
    ws2.row_dimensions[3].height = 22

    for row_idx, col in enumerate(COLUMNS, start=4):
        req_label = " (Required)" if col["required"] else " (Optional)"
        a = ws2.cell(row=row_idx, column=1)
        b = ws2.cell(row=row_idx, column=2)
        a.value = col["header"] + req_label
        b.value = col["description"] + (
            f"  |  Range: {col['min']} – {col['max']}" if "min" in col else ""
        )
        a.font = Font(bold=True, size=10, name="Calibri", color=BLUE_DARK)
        b.font = Font(size=10, name="Calibri")
        a.fill = _fill(BLUE_PALE if row_idx % 2 == 0 else WHITE)
        b.fill = _fill(BLUE_PALE if row_idx % 2 == 0 else WHITE)
        a.border = _border()
        b.border = _border()
        b.alignment = Alignment(wrap_text=True, vertical="center")
        ws2.row_dimensions[row_idx].height = 24

    # Notes
    notes_start = len(COLUMNS) + 5
    ws2.merge_cells(f"A{notes_start}:B{notes_start}")
    n = ws2.cell(row=notes_start, column=1)
    n.value = (
        "Notes:  • Columns marked * are required.  "
        "• The sample row (blue) on the Data sheet shows example values — replace or delete it before uploading.  "
        "• Each row = one vehicle's telematics for one month.  "
        "• Vehicle_Registration must exactly match the registration number in the AutoRisk system."
    )
    n.font = Font(size=9, color="64748B", italic=True, name="Calibri")
    n.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.row_dimensions[notes_start].height = 56

    return wb


# ── Parse uploaded workbook ────────────────────────────────────────────────────

def _safe(val, cast=str, default=None):
    if val is None or str(val).strip() == "":
        return default
    try:
        return cast(str(val).strip())
    except Exception:
        return default


def parse_workbook(wb: openpyxl.Workbook) -> List[dict]:
    ws = None
    for name in wb.sheetnames:
        if "data" in name.lower() or "telematics" in name.lower():
            ws = wb[name]
            break
    if ws is None:
        ws = wb.active

    # Find header row (first row with "registration" or "vehicle" in any cell)
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
        if any(str(c or "").lower().strip().startswith(("vehicle", "registration")) for c in row):
            header_row = [str(c or "").strip().lower().replace("*", "").replace(" ", "_") for c in row]
            break
    if header_row is None:
        raise ValueError("Could not find a header row. Ensure the template is used unchanged.")

    # Map header names to column indices
    col_map = {c: i for i, c in enumerate(header_row) if c}

    def get(row_vals, *keys):
        for k in keys:
            for hk, idx in col_map.items():
                if k in hk and idx < len(row_vals):
                    return row_vals[idx]
        return None

    records = []
    # Start from row after header (+1 since we skip the sample row if it looks like example data)
    data_start = None
    for row_num, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
        cells = [str(c or "").strip().lower() for c in row]
        if any(k in " ".join(cells) for k in ("vehicle_registration", "registration")):
            data_start = row_num + 1
            break

    if data_start is None:
        raise ValueError("Header row not found.")

    for row in ws.iter_rows(min_row=data_start, values_only=True):
        reg = _safe(get(row, "registration", "vehicle_reg"), str)
        if not reg:
            continue
        # Skip example rows
        if reg.lower() in ("abc1234", "example", "sample"):
            continue

        records.append({
            "registration_number":          reg,
            "customer_name":                _safe(get(row, "customer_name", "customer"), str, ""),
            "period_month":                 _safe(get(row, "period_month", "period", "month"), str,
                                                   datetime.utcnow().strftime("%Y-%m")),
            "avg_monthly_mileage_km":       max(50.0, min(5000.0, _safe(get(row, "mileage", "avg_monthly"), float, 800.0))),
            "hard_braking_events_monthly":  max(0, min(20, _safe(get(row, "hard_braking", "braking"), int, 0))),
            "speeding_incidents_monthly":   max(0, min(10, _safe(get(row, "speeding"), int, 0))),
            "night_driving_pct":            max(0.0, min(100.0, _safe(get(row, "night_driving", "night"), float, 10.0))),
            "notes":                        _safe(get(row, "notes"), str, ""),
        })

    if not records:
        raise ValueError("No data rows found. Ensure at least one data row is present.")

    return records


# ── Routes ─────────────────────────────────────────────────────────────────────

@router.get("/template")
async def download_template(_: User = Depends(get_current_user)):
    """Download the AutoRisk telematics Excel upload template."""
    wb = build_template_workbook()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"AutoRisk_Telematics_Template_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/upload")
async def upload_telematics(
    file: UploadFile = File(...),
    current: User = Depends(get_current_user),
):
    """
    Upload a filled telematics Excel file.
    Parses every data row, attempts to match vehicles by registration number,
    stores records in MongoDB, and returns all parsed rows.
    """
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Only .xlsx or .xls files are accepted.")

    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        rows = parse_workbook(wb)
    except ValueError as e:
        raise HTTPException(422, str(e))
    except Exception as e:
        raise HTTPException(422, f"Could not parse file: {e}")

    saved = []
    for row in rows:
        # Try to resolve vehicle_id / customer_id from registration number
        veh_doc = await db.vehicles.find_one(
            {"registration_number": {"$regex": f"^{row['registration_number']}$", "$options": "i"}},
            {"_id": 0, "id": 1, "customer_id": 1},
        )

        record = TelematicsRecord(
            **row,
            vehicle_id=veh_doc["id"] if veh_doc else None,
            customer_id=veh_doc["customer_id"] if veh_doc else None,
            uploaded_by=current.id,
        )
        await db.telematics.insert_one(record.model_dump(mode="json"))
        saved.append(record.model_dump(mode="json"))

    return {
        "uploaded": len(saved),
        "records": saved,
        "unmatched": [r["registration_number"] for r in saved if not r.get("vehicle_id")],
    }


@router.get("")
async def list_telematics(
    vehicle_id: Optional[str] = None,
    registration: Optional[str] = None,
    _: User = Depends(get_current_user),
):
    query = {}
    if vehicle_id:
        query["vehicle_id"] = vehicle_id
    if registration:
        query["registration_number"] = {"$regex": registration, "$options": "i"}
    docs = await db.telematics.find(query, _proj).sort("uploaded_at", -1).to_list(500)
    return docs


@router.get("/vehicle/{vehicle_id}/latest")
async def latest_for_vehicle(vehicle_id: str, _: User = Depends(get_current_user)):
    """Return the most recent telematics record for a vehicle (used by the calculator)."""
    docs = await db.telematics.find(
        {"vehicle_id": vehicle_id}, _proj
    ).sort("uploaded_at", -1).limit(1).to_list(1)
    if not docs:
        raise HTTPException(404, "No telematics record found for this vehicle")
    return docs[0]


@router.delete("/{record_id}")
async def delete_record(record_id: str, _: User = Depends(get_current_user)):
    result = await db.telematics.delete_one({"id": record_id})
    if result.deleted_count == 0:
        raise HTTPException(404, "Record not found")
    return {"message": "Deleted"}
